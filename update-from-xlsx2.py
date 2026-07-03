import openpyxl, requests, sys, re
sys.stdout.reconfigure(encoding='utf-8')

SUPABASE_URL = 'https://ycztjetxwpfgtrzeyytt.supabase.co'
SUPABASE_KEY = 'sb_secret_ZWfZCxFwF_SteBQsC8mqZA_wadaDLRd'
sb_h = {'apikey': SUPABASE_KEY, 'Authorization': f'Bearer {SUPABASE_KEY}', 'Content-Type': 'application/json'}

# --- Load Active Contracts ---
ac_wb = openpyxl.load_workbook(r'C:\Users\mmand\Downloads\Active Contracts 3.20.26.xlsx')
ac_ws = ac_wb.active
ac_headers = [c.value for c in next(ac_ws.iter_rows(min_row=1, max_row=1))]
ac_by_assoc = {}
ac_by_name = {}
for row in ac_ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(ac_headers, row))
    assoc = str(d.get('AssocID') or '').strip()
    email = (d.get('Email') or '').strip().lower()
    name = (d.get('Name') or '').strip().upper()
    if assoc and assoc != '0':
        ac_by_assoc[assoc] = {'email': email, 'name': name}
    if name:
        ac_by_name[name] = {'email': email, 'assoc': assoc}

# --- Load Producer List ---
pl_wb = openpyxl.load_workbook(r'C:\Users\mmand\Downloads\Producer List 3.20.26.xlsx')
pl_ws = pl_wb.active
pl_headers = [c.value for c in next(pl_ws.iter_rows(min_row=1, max_row=1))]
pl_by_assoc = {}
for row in pl_ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(pl_headers, row))
    assoc = str(d.get('Associate ID') or '').strip()
    email = (d.get('Company Email') or '').strip().lower()
    states_raw = (d.get('Life-and-Health Licensed States') or '').strip()
    name = (d.get('Agent') or '').strip().upper()
    if assoc and assoc != '0':
        pl_by_assoc[assoc] = {'email': email, 'states': states_raw, 'name': name}

print(f'Active Contracts: {len(ac_by_assoc)} | Producer List: {len(pl_by_assoc)}')

# --- Load ALL customers with pagination ---
customers = []
offset = 0
while True:
    resp = requests.get(
        f'{SUPABASE_URL}/rest/v1/customers?select=id,agent_name,associate_id,company_email,states&limit=1000&offset={offset}',
        headers=sb_h)
    batch = resp.json()
    if not batch:
        break
    customers.extend(batch)
    offset += len(batch)
    if len(batch) < 1000:
        break

print(f'Total customers: {len(customers)}\n')

updated_email = 0
updated_states = 0

for cust in customers:
    cid = cust['id']
    assoc = str(cust.get('associate_id') or '').strip()
    current_email = (cust.get('company_email') or '').strip().lower()
    current_states = cust.get('states')
    name = (cust.get('agent_name') or '').strip().upper()

    patch = {}

    src = pl_by_assoc.get(assoc) or ac_by_assoc.get(assoc)
    if not src and name:
        src = ac_by_name.get(name)

    if not src:
        continue

    # Email - only update if missing
    new_email = src.get('email', '')
    if not current_email and new_email:
        patch['company_email'] = new_email

    # States - only update if missing, from producer list
    pl_src = pl_by_assoc.get(assoc)
    if pl_src and not current_states:
        states_raw = pl_src.get('states', '')
        if states_raw:
            patch['states'] = ', '.join(s.strip() for s in states_raw.split(',') if s.strip())

    if patch:
        r = requests.patch(f'{SUPABASE_URL}/rest/v1/customers?id=eq.{cid}', headers=sb_h, json=patch)
        if r.status_code == 204:
            parts = []
            if 'company_email' in patch:
                parts.append(f'email={patch["company_email"]}')
                updated_email += 1
            if 'states' in patch:
                parts.append(f'states({len(patch["states"].split(","))} states)')
                updated_states += 1
            print(f'  ✅ {cust["agent_name"]} -> {", ".join(parts)}')
        else:
            print(f'  ❌ {cust["agent_name"]}: {r.status_code} {r.text[:80]}')

print(f'\nDone. Emails updated: {updated_email} | States updated: {updated_states}')
