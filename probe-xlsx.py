import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

for fname, label in [
    (r'C:\Users\mmand\Downloads\Active Contracts 3.20.26.xlsx', 'Active Contracts'),
    (r'C:\Users\mmand\Downloads\Producer List 3.20.26.xlsx', 'Producer List'),
]:
    print(f'\n=== {label} ===')
    wb = openpyxl.load_workbook(fname)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    print('Headers:', headers)
    for row in ws.iter_rows(min_row=2, max_row=4, values_only=True):
        print(row)
